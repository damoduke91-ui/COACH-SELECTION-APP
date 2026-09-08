"""Safely sync the latest AFL Players_Flat workbook into Supabase.

The importer resolves the controlled active season from app_settings and
competition_seasons. It never relies on a calendar-year guess, and every read,
upsert, verification, and stale-row deletion is scoped to environment, season,
round, and AFL club.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
import urllib.error
import urllib.parse
import urllib.request
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook


DEFAULT_BASE_DIR = Path(r"C:\AFL_Teams")
MIN_VALID_PLAYER_ROWS = 40
MIN_VALID_TEAM_ROWS = 20
VALID_ENVIRONMENTS = {"production", "preview"}


def clean_text(value: object) -> str:
    return str(value or "").strip()


def load_env_file(path: Path) -> dict[str, str]:
    values: dict[str, str] = {}
    if not path.is_file():
        raise RuntimeError(f"Supabase environment file was not found: {path}")

    for raw_line in path.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        values[key.strip()] = value.strip().strip('"').strip("'")
    return values


def select_workbook_path(base_dir: Path) -> Path:
    current_dir = base_dir / "Current"
    candidates = [
        path
        for path in (
            current_dir / "AFL_Teams_Latest.xlsx",
            current_dir / "AFL_Teams_Latest.pending.xlsx",
        )
        if path.is_file()
    ]
    if not candidates:
        raise RuntimeError(f"No current AFL teams workbook was found in {current_dir}.")
    return max(candidates, key=lambda path: path.stat().st_mtime)


def read_workbook(path: Path) -> tuple[int, list[dict[str, object]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if "Summary" not in workbook.sheetnames or "Players_Flat" not in workbook.sheetnames:
            raise RuntimeError("Workbook must contain Summary and Players_Flat sheets.")

        summary = {
            clean_text(row[0]): clean_text(row[1])
            for row in workbook["Summary"].iter_rows(values_only=True)
            if len(row) >= 2 and clean_text(row[0])
        }
        source_url = summary.get("source_url", "")
        round_match = re.search(r"CD_R\d{7}(\d{2})(?:\D|$)", source_url, re.IGNORECASE)
        if not round_match:
            raise RuntimeError(f"Could not detect the AFL round from source_url: {source_url}")

        round_number = int(round_match.group(1))
        if not 1 <= round_number <= 24:
            raise RuntimeError(f"Detected AFL round is outside the expected range: {round_number}")

        values = workbook["Players_Flat"].iter_rows(values_only=True)
        try:
            headers = [clean_text(value) for value in next(values)]
        except StopIteration as exc:
            raise RuntimeError("Players_Flat is empty.") from exc

        required = {"player_name", "team", "section", "line_group", "raw_position"}
        missing = required - set(headers)
        if missing:
            raise RuntimeError(f"Players_Flat is missing columns: {', '.join(sorted(missing))}")

        rows: list[dict[str, object]] = []
        for row_values in values:
            source = {
                headers[index]: row_values[index] if index < len(row_values) else ""
                for index in range(len(headers))
                if headers[index]
            }
            player_name = clean_text(source.get("player_name"))
            afl_team = clean_text(source.get("team"))
            if not player_name or not afl_team:
                continue

            section = clean_text(source.get("section")).lower()
            line_group = clean_text(source.get("line_group")).lower()
            raw_position = clean_text(source.get("raw_position"))
            role_marker = " ".join((section, line_group, raw_position.lower()))
            role2 = (
                "EMERG"
                if "emerg" in role_marker
                else "INT"
                if "interchange" in role_marker or raw_position.upper() == "INT"
                else ""
            )
            rows.append(
                {
                    "round": round_number,
                    "player_name": player_name,
                    "afl_team": afl_team,
                    "role1": raw_position,
                    "role2": role2,
                }
            )
    finally:
        workbook.close()

    if len(rows) < MIN_VALID_PLAYER_ROWS:
        raise RuntimeError(
            f"Workbook has only {len(rows)} player rows; at least {MIN_VALID_PLAYER_ROWS} are required."
        )

    teams: dict[str, list[dict[str, object]]] = defaultdict(list)
    for row in rows:
        teams[str(row["afl_team"])].append(row)
    if len(teams) < 2:
        raise RuntimeError("Workbook contains fewer than two announced teams.")

    undersized = {
        team: len(team_rows)
        for team, team_rows in teams.items()
        if len(team_rows) < MIN_VALID_TEAM_ROWS
    }
    if undersized:
        raise RuntimeError(f"Refusing partial team data with small rosters: {undersized}")

    names = [str(row["player_name"]) for row in rows]
    duplicate_names = sorted({name for name in names if names.count(name) > 1})
    if duplicate_names:
        raise RuntimeError(
            f"Duplicate player names detected in round {round_number}: {', '.join(duplicate_names)}"
        )
    return round_number, rows


class SupabaseRest:
    def __init__(self, url: str, service_key: str):
        self.rest_url = url.rstrip("/") + "/rest/v1"
        self.headers = {
            "apikey": service_key,
            "Authorization": f"Bearer {service_key}",
            "Content-Type": "application/json",
        }

    def request(self, table: str, method: str, query: str, payload=None, prefer: str | None = None):
        headers = dict(self.headers)
        if prefer:
            headers["Prefer"] = prefer
        data = None if payload is None else json.dumps(payload).encode("utf-8")
        request = urllib.request.Request(
            self.rest_url + "/" + table + "?" + query,
            data=data,
            headers=headers,
            method=method,
        )
        try:
            with urllib.request.urlopen(request, timeout=60) as response:
                body = response.read().decode("utf-8")
                return json.loads(body) if body else None
        except urllib.error.HTTPError as error:
            body = error.read().decode("utf-8", errors="replace")
            raise RuntimeError(f"Supabase {method} failed: HTTP {error.code}: {body}") from error

    def controlled_season(self, environment: str) -> int:
        settings_query = urllib.parse.urlencode(
            {"select": "environment,season_year", "environment": f"eq.{environment}"}
        )
        settings_rows = self.request("app_settings", "GET", settings_query) or []
        if len(settings_rows) != 1:
            raise RuntimeError(f"Expected one app_settings row for {environment}.")
        season_year = int(settings_rows[0]["season_year"])

        season_query = urllib.parse.urlencode(
            {
                "select": "environment,season_year,status",
                "environment": f"eq.{environment}",
                "season_year": f"eq.{season_year}",
            }
        )
        season_rows = self.request("competition_seasons", "GET", season_query) or []
        if len(season_rows) != 1 or season_rows[0].get("status") != "active":
            raise RuntimeError(
                f"Season {season_year} for {environment} is not the single controlled active season."
            )
        return season_year

    def team_rows(self, environment: str, season_year: int, round_number: int, team: str):
        query = urllib.parse.urlencode(
            {
                "select": "id,player_name,afl_team,role1,role2",
                "environment": f"eq.{environment}",
                "season_year": f"eq.{season_year}",
                "round": f"eq.{round_number}",
                "afl_team": f"eq.{team}",
            }
        )
        return self.request("weekly_team_lists", "GET", query) or []

    def upsert(self, rows: list[dict[str, object]]) -> None:
        self.request(
            "weekly_team_lists",
            "POST",
            "on_conflict=environment,season_year,round,player_name",
            rows,
            "resolution=merge-duplicates",
        )

    def delete_id(self, row_id: str) -> None:
        self.request("weekly_team_lists", "DELETE", urllib.parse.urlencode({"id": f"eq.{row_id}"}))

    def patch_settings(self, environment: str, payload: dict[str, object]) -> None:
        query = urllib.parse.urlencode({"environment": f"eq.{environment}"})
        updated = self.request("app_settings", "PATCH", query, payload, "return=representation") or []
        if len(updated) != 1:
            raise RuntimeError(f"Could not update app_settings for {environment}.")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--base-dir", type=Path, default=DEFAULT_BASE_DIR)
    parser.add_argument("--environment", choices=sorted(VALID_ENVIRONMENTS), default="production")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    base_dir: Path = args.base_dir
    environment: str = args.environment
    workbook_path = select_workbook_path(base_dir)
    round_number, rows = read_workbook(workbook_path)

    env = load_env_file(base_dir / ".env")
    supabase_url = env.get("NEXT_PUBLIC_SUPABASE_URL") or env.get("SUPABASE_URL") or os.environ.get("SUPABASE_URL", "")
    service_key = env.get("SUPABASE_SERVICE_ROLE_KEY") or os.environ.get("SUPABASE_SERVICE_ROLE_KEY", "")
    if not supabase_url or not service_key:
        raise RuntimeError("Supabase URL or service-role key is missing.")

    client = SupabaseRest(supabase_url, service_key)
    season_year = client.controlled_season(environment)
    scoped_rows = [
        {**row, "environment": environment, "season_year": season_year}
        for row in rows
    ]
    teams: dict[str, list[dict[str, object]]] = defaultdict(list)
    for row in scoped_rows:
        teams[str(row["afl_team"])].append(row)

    client.patch_settings(
        environment,
        {
            "team_list_sync_status": "running",
            "team_list_sync_at": datetime.now(timezone.utc).isoformat(),
            "team_list_sync_round": round_number,
            "team_list_sync_player_count": len(scoped_rows),
            "team_list_sync_team_count": len(teams),
            "team_list_sync_message": f"AFL {season_year} Round {round_number} is being synced.",
        },
    )
    print(
        f"Syncing {environment} AFL {season_year} Round {round_number}: "
        f"{len(scoped_rows)} rows across {len(teams)} announced teams."
    )

    for team in sorted(teams):
        incoming = teams[team]
        incoming_names = {str(row["player_name"]) for row in incoming}
        existing = client.team_rows(environment, season_year, round_number, team)
        client.upsert(incoming)
        for stale_row in existing:
            if stale_row.get("player_name") not in incoming_names:
                client.delete_id(str(stale_row["id"]))

        verified = client.team_rows(environment, season_year, round_number, team)
        verified_names = {str(row.get("player_name")) for row in verified}
        if verified_names != incoming_names:
            raise RuntimeError(
                f"Verification failed for {team}; "
                f"missing={sorted(incoming_names - verified_names)}, "
                f"extra={sorted(verified_names - incoming_names)}"
            )
        print(f"Verified {team}: {len(verified)} rows.")

    client.patch_settings(
        environment,
        {
            "latest_team_list_round": round_number,
            "team_list_sync_status": "success",
            "team_list_sync_at": datetime.now(timezone.utc).isoformat(),
            "team_list_sync_round": round_number,
            "team_list_sync_player_count": len(scoped_rows),
            "team_list_sync_team_count": len(teams),
            "team_list_sync_message": f"AFL {season_year} Round {round_number} synced successfully.",
        },
    )
    print(f"AFL {season_year} Round {round_number} sync completed successfully.")


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        sys.exit(1)
